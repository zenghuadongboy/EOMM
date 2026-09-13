import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="D3D3D3"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="D3D3D3"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def build_docx():
    filename = "d:/大学/基于.NET的企业运维管理平/RESTful API接口设计与数据库汇总表.docx"
    doc = docx.Document()

    # 页边距
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 默认字体
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # 文档大标题
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("基于 .NET 的企业运维管理平台 - 架构与接口设计文档")
    run_title.font.name = "黑体"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("包含数据库表设计汇总表 (表3.1) 与 典型 RESTful API 接口设计表 (表3.2)")
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_paragraph() # 空行

    # --- 1. 表3.1 数据库表 ---
    p_t1_title = doc.add_paragraph()
    p_t1_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_t1_title = p_t1_title.add_run("表3.1 数据库表")
    r_t1_title.font.name = "黑体"
    r_t1_title.font.size = Pt(11)
    r_t1_title.font.bold = True

    db_data = [
        ("sys_users", "Id、UserName、PasswordHash、RealName、Email、Status、CreateTime", "保存用户身份账号、密码哈希、邮箱及账号状态信息"),
        ("sys_roles", "Id、RoleCode、RoleName、Description、CreateTime", "保存系统角色编码、角色名称及说明信息"),
        ("sys_user_roles", "UserId、RoleId、CreateTime", "建立用户与角色的多对多映射关联"),
        ("sys_permissions", "Id、PermissionCode、PermissionName、Path、Method、Description、CreateTime", "保存权限编码、权限名称、路由路径及请求方法"),
        ("sys_role_permissions", "Id、RoleId、PermissionId", "建立角色与权限的多对多授权关联"),
        ("asset_types", "Id、TypeCode、TypeName、Description、CreateTime", "保存设备资产分类及类型编码信息"),
        ("assets", "Id、AssetNo、AssetName、AssetTypeId、IpAddress、Location、Status、Description、CreateTime", "保存设备资产编号、名称、类型关联、IP地址、部署位置及运行状态"),
        ("alarm_roles", "Id、RuleName、AssetTypeId、Metric、Threshold、Severity、Status、Description、CreateTime", "保存告警判断规则名称、关联设备类型、监控指标、阈值及告警级别"),
        ("alarm_records", "Id、AlarmNo、AssetId、RuleId、Content、Severity、Status、TriggeredAt、Handler、ResolutionNotes、ResolvedAt", "保存触发的设备告警记录、告警等级、处理人及解决状态"),
        ("work_orders", "Id、WorkOrderNo、Source、Priority、AssetId、AlarmId、Handler、Status、Title、Description、CreateTime、FinishedAt", "保存运维处置工单编号、关联告警/设备、优先级、负责人及完工状态"),
        ("work_order_logs", "Id、WorkOrderId、OperatorId、OperatorName、PreviousStatus、NewStatus、ActionNote、CreateTime", "保存工单状态流转变迁履历、操作人及处理备注"),
        ("operation_logs", "Id、Username、Module、Action、IpAddress、RequestParam、Success、ExecutionTimeMs、Description、CreatedAt", "保存用户登录及敏感业务操作审计日志"),
        ("system_logs", "Id、Level、Source、Message、ExceptionDetails、CreatedAt", "保存后台自动化告警引擎运行状态及系统异常日志")
    ]

    table1 = doc.add_table(rows=1, cols=3)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table1)

    hdr1_cells = table1.rows[0].cells
    hdr1_titles = ["表名", "主要字段", "作用"]
    col1_widths = [Inches(1.5), Inches(3.2), Inches(2.2)]

    for i, h_text in enumerate(hdr1_titles):
        hdr1_cells[i].width = col1_widths[i]
        set_cell_background(hdr1_cells[i], "1F497D")
        set_cell_margins(hdr1_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr1_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.name = "黑体"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for row_idx, (tbl_name, fields, purpose) in enumerate(db_data):
        row_cells = table1.add_row().cells
        bg_color = "F2F5F9" if row_idx % 2 == 1 else "FFFFFF"

        row_cells[0].width = col1_widths[0]
        set_cell_background(row_cells[0], bg_color)
        set_cell_margins(row_cells[0], top=80, bottom=80, left=100, right=100)
        p0 = row_cells[0].paragraphs[0]
        r0 = p0.add_run(tbl_name)
        r0.font.name = "Consolas"
        r0.font.size = Pt(9.0)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

        row_cells[1].width = col1_widths[1]
        set_cell_background(row_cells[1], bg_color)
        set_cell_margins(row_cells[1], top=80, bottom=80, left=100, right=100)
        p1 = row_cells[1].paragraphs[0]
        r1 = p1.add_run(fields)
        r1.font.name = "Consolas"
        r1.font.size = Pt(8.5)

        row_cells[2].width = col1_widths[2]
        set_cell_background(row_cells[2], bg_color)
        set_cell_margins(row_cells[2], top=80, bottom=80, left=100, right=100)
        p2 = row_cells[2].paragraphs[0]
        r2 = p2.add_run(purpose)
        r2.font.name = "宋体"
        r2.font.size = Pt(9.0)

    doc.add_paragraph() # 分隔空行
    doc.add_paragraph()

    # --- 2. 表3.2 典型 RESTful API 接口设计表 ---
    p_t2_title = doc.add_paragraph()
    p_t2_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_t2_title = p_t2_title.add_run("表3.2 系统典型 RESTful API 接口设计表")
    r_t2_title.font.name = "黑体"
    r_t2_title.font.size = Pt(11)
    r_t2_title.font.bold = True

    api_data = [
        ("身份认证", "/api/auth/login", "POST", "用户登录鉴权并生成 JWT Token", "username, password"),
        ("身份认证", "/api/auth/me", "GET", "获取当前登录用户的基本信息与权限列表", "Bearer Token (Header)"),
        ("身份认证", "/api/auth/change-password", "POST", "修改当前登录用户的密码", "oldPassword, newPassword"),
        ("设备资产", "/api/asset/list", "GET", "多条件筛选与分页查询资产台账", "keyword, type, status, pageIndex, pageSize"),
        ("设备资产", "/api/asset/{id}", "GET", "根据 ID 查询指定设备资产详细档案", "id (路径参数)"),
        ("设备资产", "/api/asset/save", "POST", "新建或修改设备资产档案信息", "assetNo, name, assetTypeId, ipAddress, location"),
        ("设备资产", "/api/asset/{id}", "DELETE", "根据 ID 删除指定设备资产记录", "id (路径参数)"),
        ("告警管理", "/api/alarm/records", "GET", "分页查询自动化触发的告警记录列表", "status, severity, pageIndex, pageSize"),
        ("告警管理", "/api/alarm/process", "POST", "处置闭环告警（结案并填写复盘说明）", "alarmId, status, resolutionNotes"),
        ("告警管理", "/api/alarm/rules", "GET", "获取系统配置的告警规则列表", "无"),
        ("告警管理", "/api/alarm/rule/save", "POST", "新增或编辑告警触发条件规则", "ruleName, assetTypeId, metric, threshold, severity"),
        ("运维工单", "/api/workorder/list", "GET", "分页与条件查询运维处置工单", "keyword, status, priority, pageIndex, pageSize"),
        ("运维工单", "/api/workorder/save", "POST", "创建或编辑运维例行派发工单", "title, assetId, handler, priority, description"),
        ("运维工单", "/api/workorder/create-from-alarm", "POST", "告警一键联动派发运维处置工单", "alarmRecordId, handler"),
        ("运维工单", "/api/workorder/change-status", "POST", "工单状态机流转（接单/处理/完工/关单）", "workOrderId, newStatus, note, resolutionNotes"),
        ("运维工单", "/api/workorder/logs/{id}", "GET", "查询指定工单的完整流转履历日志", "id (工单ID)"),
        ("用户角色", "/api/user/list", "GET", "获取系统用户列表及绑定的角色信息", "无"),
        ("用户角色", "/api/user/save", "POST", "保存/更新用户账号及分配的角色", "username, realName, email, roleId, isEnabled"),
        ("用户角色", "/api/role/list", "GET", "获取角色列表及关联的权限列表", "无"),
        ("用户角色", "/api/role/assign-permissions", "POST", "为指定角色配置授权权限点集合", "roleId, permissionIds"),
        ("大屏大盘", "/api/dashboard/overview", "GET", "运维大屏聚合数据（资产分布、告警趋势、Top5故障）", "无"),
        ("审计日志", "/api/log/operation", "GET", "分页查询关键业务敏感操作审计日志", "username, module, pageIndex, pageSize")
    ]

    table2 = doc.add_table(rows=1, cols=5)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table2)

    hdr2_cells = table2.rows[0].cells
    hdr2_titles = ["模块名称", "请求路径 (URI)", "HTTP方法", "接口功能说明", "典型入参 / 格式"]
    col2_widths = [Inches(1.0), Inches(1.8), Inches(0.8), Inches(1.8), Inches(1.5)]

    for i, h_text in enumerate(hdr2_titles):
        hdr2_cells[i].width = col2_widths[i]
        set_cell_background(hdr2_cells[i], "2F5597")
        set_cell_margins(hdr2_cells[i], top=100, bottom=100, left=100, right=100)
        p = hdr2_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.name = "黑体"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for row_idx, (mod_name, uri, method, desc, params) in enumerate(api_data):
        row_cells = table2.add_row().cells
        bg_color = "F2F5F9" if row_idx % 2 == 1 else "FFFFFF"

        # Col 0: 模块名称
        row_cells[0].width = col2_widths[0]
        set_cell_background(row_cells[0], bg_color)
        set_cell_margins(row_cells[0], top=80, bottom=80, left=80, right=80)
        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(mod_name)
        r0.font.name = "宋体"
        r0.font.size = Pt(9.0)
        r0.font.bold = True

        # Col 1: URI
        row_cells[1].width = col2_widths[1]
        set_cell_background(row_cells[1], bg_color)
        set_cell_margins(row_cells[1], top=80, bottom=80, left=80, right=80)
        p1 = row_cells[1].paragraphs[0]
        r1 = p1.add_run(uri)
        r1.font.name = "Consolas"
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

        # Col 2: HTTP方法
        row_cells[2].width = col2_widths[2]
        set_cell_background(row_cells[2], bg_color)
        set_cell_margins(row_cells[2], top=80, bottom=80, left=80, right=80)
        p2 = row_cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(method)
        r2.font.name = "Consolas"
        r2.font.size = Pt(8.5)
        r2.font.bold = True
        if method == "GET":
            r2.font.color.rgb = RGBColor(0x28, 0xa7, 0x45)
        elif method == "POST":
            r2.font.color.rgb = RGBColor(0x00, 0x56, 0xb3)
        elif method == "DELETE":
            r2.font.color.rgb = RGBColor(0xd9, 0x53, 0x4f)

        # Col 3: 功能说明
        row_cells[3].width = col2_widths[3]
        set_cell_background(row_cells[3], bg_color)
        set_cell_margins(row_cells[3], top=80, bottom=80, left=80, right=80)
        p3 = row_cells[3].paragraphs[0]
        r3 = p3.add_run(desc)
        r3.font.name = "宋体"
        r3.font.size = Pt(8.5)

        # Col 4: 典型入参
        row_cells[4].width = col2_widths[4]
        set_cell_background(row_cells[4], bg_color)
        set_cell_margins(row_cells[4], top=80, bottom=80, left=80, right=80)
        p4 = row_cells[4].paragraphs[0]
        r4 = p4.add_run(params)
        r4.font.name = "Consolas"
        r4.font.size = Pt(8.0)

    doc.save(filename)
    print(f"完整 RESTful API 与数据库设计 Word 文档已导出成功: {filename}")

if __name__ == "__main__":
    build_docx()
