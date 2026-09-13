import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="D3D3D3"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="D3D3D3"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def build_docx(filename):
    doc = docx.Document()

    # 页边距
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # 默认字体
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # 文档大标题
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("基于 .NET 的企业运维管理平台 - 数据库设计")
    run_title.font.name = "黑体"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("数据库表与核心字段汇总说明表")
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_paragraph() # 空行

    # 表标题
    p_table_title = doc.add_paragraph()
    p_table_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_table_title = p_table_title.add_run("表3.1 数据库表")
    run_table_title.font.name = "黑体"
    run_table_title.font.size = Pt(11)
    run_table_title.font.bold = True

    # 表格数据
    data = [
        ("sys_users", "Id、UserName、PasswordHash、RealName、Email、Status、CreateTime", "保存用户身份账号、密码哈希、邮箱及账号状态信息"),
        ("sys_roles", "Id、RoleCode、RoleName、Description、CreateTime", "保存系统角色编码、角色名称及说明信息"),
        ("sys_user_roles", "UserId、RoleId、CreateTime", "建立用户与角色的多对多映射关联"),
        ("sys_permissions", "Id、PermissionCode、PermissionName、Path、Method、Description、CreateTime", "保存权限编码、权限名称、路由路径及请求方法"),
        ("sys_role_permissions", "Id、RoleId、PermissionId", "建立角色与权限的多对多授权关联"),
        ("asset_types", "Id、TypeCode、TypeName、Description、CreateTime", "保存设备资产分类及类型编码信息"),
        ("assets", "Id、AssetNo、AssetName、AssetTypeId、IpAddress、Location、Status、Description、CreateTime", "保存设备资产编号、名称、类型关联、IP地址、部署位置及运行状态"),
        ("alarm_roles", "Id、RuleName、AssetTypeId、Metric、Threshold、Severity、Status、Description、CreateTime", "保存告警判断规则名称、关联设备类型、监控指标、阈值及告警级别"),
        ("alarm_records", "Id、AlarmNo、AssetId、RuleId、Content、Severity、Status、TriggeredAt、Handler、ResolutionNotes、ResolvedAt", "保存触发的设备告警记录、告警等级、处理人及解决状态"),
        ("work_orders", "Id、WorkOrderNo、Source、Priority、AssetId、AlarmId、Handler、Status、Title、Description、CreateTime、FinishedAt", "保存运维处置工单编号、关联告警/设备、优先级、负责人及完工状态"),
        ("work_order_logs", "Id、WorkOrderId、OperatorId、OperatorName、PreviousStatus、NewStatus、ActionNote、CreateTime", "保存工单状态流转变迁履历、操作人及处理备注"),
        ("operation_logs", "Id、Username、Module、Action、IpAddress、RequestParam、Success、ExecutionTimeMs、Description、CreatedAt", "保存用户登录及敏感业务操作审计日志"),
        ("system_logs", "Id、Level、Source、Message、ExceptionDetails、CreatedAt", "保存后台自动化告警引擎运行状态及系统异常日志")
    ]

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # 表头
    hdr_cells = table.rows[0].cells
    headers = ["表名", "主要字段", "作用"]
    col_widths = [Inches(1.5), Inches(3.0), Inches(2.3)]

    for i, header_text in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "1F497D")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header_text)
        run.font.name = "黑体"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # 表格内容填充
    for row_idx, (tbl_name, fields, purpose) in enumerate(data):
        row_cells = table.add_row().cells
        bg_color = "F2F5F9" if row_idx % 2 == 1 else "FFFFFF"

        # Col 0: 表名
        row_cells[0].width = col_widths[0]
        set_cell_background(row_cells[0], bg_color)
        set_cell_margins(row_cells[0], top=100, bottom=100, left=120, right=120)
        p0 = row_cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run(tbl_name)
        r0.font.name = "Consolas"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

        # Col 1: 主要字段
        row_cells[1].width = col_widths[1]
        set_cell_background(row_cells[1], bg_color)
        set_cell_margins(row_cells[1], top=100, bottom=100, left=120, right=120)
        p1 = row_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(fields)
        r1.font.name = "Consolas"
        r1.font.size = Pt(9.0)

        # Col 2: 作用
        row_cells[2].width = col_widths[2]
        set_cell_background(row_cells[2], bg_color)
        set_cell_margins(row_cells[2], top=100, bottom=100, left=120, right=120)
        p2 = row_cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r2 = p2.add_run(purpose)
        r2.font.name = "宋体"
        r2.font.size = Pt(9.5)

    doc.save(filename)
    print(f"Word 文档已成功生成: {filename}")

if __name__ == "__main__":
    build_docx("d:/大学/基于.NET的企业运维管理平/数据库表设计汇总表.docx")
